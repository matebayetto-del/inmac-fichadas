
# -*- coding: utf-8 -*-
"""
INMAC · Motor universal de fichajes para planillas base TORNU / Hemoterapia / Cubiertas.

Versión v4 final:
- Detecta planilla base y archivo de fichajes sin depender del nombre.
- Detecta el mes desde el archivo Original del fichero.
- Reubica las fórmulas diarias usando las fórmulas originales de la plantilla base:
  * laboral lunes-viernes
  * sábado
  * domingo / feriado
- Carga entradas y salidas como horas reales de Excel.
- Marca AUSENTE en días laborables sin fichada.
- Actualiza fórmulas de resumen: AUSENCIA, ENFERMEDAD, VIANDA, ART, FERIADO, VACACIONES.
- Corrige el formato visual de todas las celdas de resumen para mostrar enteros en negro.
- Resalta nombres en escala roja según faltas y legajos en amarillo cuando el vínculo es dudoso.
- Limpia alertas anteriores antes de recalcularlas.
- No modifica datos manuales salvo celdas de carga diaria, fórmulas de resumen y alertas visuales.
"""

from __future__ import annotations

import os
import re
import copy
import unicodedata
import calendar
from pathlib import Path
from datetime import datetime, date, time
from collections import defaultdict
from difflib import SequenceMatcher

import xlrd
import holidays
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font, PatternFill

PAIS_FERIADOS = "AR"
SUBDIV_FERIADOS = None
UMBRAL_NOMBRE = 0.86
UMBRAL_NOMBRE_SIMILAR = 0.90

# Alertas visuales.
UMBRAL_FALTAS_ROJO = 3
UMBRAL_IGNORADO_CERCANO = 0.70

ESCALA_ROJO_FALTAS = [
    (11, "E06666", "ROJO FUERTE"),
    (8, "EA9999", "ROJO MEDIO"),
    (5, "F4CCCC", "ROJO SUAVE"),
    (3, "FCE8E6", "ROJO MUY SUAVE"),
]

COLOR_AMARILLO_DUDOSO = "FFF2CC"
COLOR_TEXTO_RESUMEN = "000000"

CORRECCIONES_NOMBRE = {
    "surez": "suarez",
    "suárez": "suarez",
    "bayeto": "bayetto",
    "calvatti": "calvitti",
    "alexiz": "alexis",
    "davalos": "davalo",
    "javer": "javier",
    "ramura": "ramua",
    "jimenz": "jimenez",
    "francico": "francisco",
    "reyes": "reyez",
}

HOJAS_AUXILIARES = [
    "VINCULOS_CARGADOS",
    "IGNORADOS_FICHADOR",
    "DIAGNOSTICO",
    "ALERTAS_PROCESO",
]


# ---------------------------------------------------------------------------
# Normalización / matching
# ---------------------------------------------------------------------------

def quitar_tildes(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    )


def normalizar_texto(valor) -> str:
    texto = "" if valor is None else str(valor)
    texto = quitar_tildes(texto).lower()
    texto = re.sub(r"[^a-z0-9]+", " ", texto).strip()
    tokens = [CORRECCIONES_NOMBRE.get(t, t) for t in texto.split()]
    return " ".join(tokens)


def normalizar_header(valor) -> str:
    return normalizar_texto(valor).replace("%", "").strip()


def clave_nombre(valor) -> str:
    return " ".join(sorted(normalizar_texto(valor).split()))


def puntaje_nombre(a, b) -> float:
    ca, cb = clave_nombre(a), clave_nombre(b)
    if not ca or not cb:
        return 0.0
    if ca == cb:
        return 1.0
    ta, tb = set(ca.split()), set(cb.split())
    cobertura = len(ta & tb) / max(1, min(len(ta), len(tb)))
    return max(SequenceMatcher(None, ca, cb).ratio(), cobertura * 0.95)


def nombres_compatibles(a, b) -> bool:
    ca, cb = clave_nombre(a), clave_nombre(b)
    if ca and cb and ca == cb:
        return True
    ta, tb = set(normalizar_texto(a).split()), set(normalizar_texto(b).split())
    if ta and tb and (ta.issubset(tb) or tb.issubset(ta)):
        return True
    return puntaje_nombre(a, b) >= UMBRAL_NOMBRE


def normalizar_legajo(valor) -> str | None:
    if valor is None:
        return None
    txt = str(valor).replace(".0", "").strip()
    dig = re.sub(r"\D", "", txt)
    if not dig:
        return None
    return dig.lstrip("0") or "0"


def derivar_legajo_desde_id(valor) -> str | None:
    if valor is None:
        return None
    dig = re.sub(r"\D", "", str(valor).replace(".0", ""))
    if not dig:
        return None
    if len(dig) > 3 and dig[:3] in {"110", "130"}:
        dig = dig[3:]
    return dig.lstrip("0") or "0"


# ---------------------------------------------------------------------------
# Lectura de fichajes
# ---------------------------------------------------------------------------

def leer_fichajes(path: str | Path) -> list[dict]:
    libro = xlrd.open_workbook(str(path))
    hoja = libro.sheet_by_name("Original") if "Original" in libro.sheet_names() else libro.sheet_by_index(0)

    headers = {}
    for c in range(hoja.ncols):
        headers[normalizar_header(hoja.cell_value(0, c))] = c

    def col(*opciones):
        for op in opciones:
            key = normalizar_header(op)
            if key in headers:
                return headers[key]
        raise ValueError(f"No encontré columna entre: {opciones}")

    c_id = col("ID")
    c_nombre = col("Nombre")
    c_hora = col("Hora", "Fecha")

    registros = []
    for r in range(1, hoja.nrows):
        raw_id = str(hoja.cell_value(r, c_id)).replace(".0", "").strip()
        nombre = str(hoja.cell_value(r, c_nombre)).strip()
        hora_raw = hoja.cell_value(r, c_hora)
        if not raw_id and not nombre and not hora_raw:
            continue

        dt = None
        if isinstance(hora_raw, (float, int)):
            try:
                dt = datetime(*xlrd.xldate_as_tuple(hora_raw, libro.datemode))
            except Exception:
                dt = None
        else:
            for fmt in (
                "%Y/%m/%d %H:%M:%S",
                "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y %H:%M:%S",
                "%d/%m/%Y %H:%M",
            ):
                try:
                    dt = datetime.strptime(str(hora_raw).strip(), fmt)
                    break
                except Exception:
                    pass

        if dt is None:
            continue

        registros.append({
            "id_raw": raw_id,
            "legajo": derivar_legajo_desde_id(raw_id),
            "nombre": nombre,
            "nombre_norm": normalizar_texto(nombre),
            "dt": dt,
        })

    if not registros:
        raise ValueError("El archivo de fichajes no tiene registros válidos.")
    return registros


def detectar_mes(registros: list[dict]) -> tuple[int, int]:
    conteo = defaultdict(int)
    for r in registros:
        conteo[(r["dt"].year, r["dt"].month)] += 1
    return max(conteo.items(), key=lambda x: x[1])[0]


# ---------------------------------------------------------------------------
# Detección de planilla
# ---------------------------------------------------------------------------

def detectar_hoja_planilla(wb):
    mejores = []
    for ws in wb.worksheets:
        texto = []
        for r in range(1, min(ws.max_row, 5) + 1):
            for c in range(1, min(ws.max_column, 260) + 1):
                texto.append(str(ws.cell(r, c).value or ""))
        flat = normalizar_texto(" ".join(texto))
        score = 0
        if "apellido y nombre" in flat:
            score += 5
        if "entrada" in flat and "salida" in flat and "ausente" in flat:
            score += 5
        if "horas primera quincena" in flat or "horas segunda quincena" in flat:
            score += 3
        if score:
            mejores.append((score, ws.title))
    if not mejores:
        raise ValueError("No encontré una hoja de planilla compatible.")
    mejores.sort(reverse=True)
    return wb[mejores[0][1]]


def detectar_columnas_identidad(ws):
    fila_header = 1
    col_centro = col_cat = col_leg = col_nombre = col_ingreso = None
    for c in range(1, min(ws.max_column, 30) + 1):
        h = normalizar_header(ws.cell(fila_header, c).value)
        if h == "centro de costo":
            col_centro = c
        elif h == "categoria":
            col_cat = c
        elif h in {"leg", "legajo"}:
            col_leg = c
        elif "apellido y nombre" in h or "nombre y apellido" in h:
            col_nombre = c
        elif h == "ingreso":
            col_ingreso = c
    if col_leg is None or col_nombre is None:
        raise ValueError("No encontré columnas LEG y APELLIDO Y NOMBRE.")
    return {
        "fila_header": fila_header,
        "col_centro": col_centro,
        "col_categoria": col_cat,
        "col_legajo": col_leg,
        "col_nombre": col_nombre,
        "col_ingreso": col_ingreso,
    }


def detectar_bloques_diarios(ws) -> list[dict]:
    bloques = []
    for c in range(1, ws.max_column - 5):
        vals = [normalizar_header(ws.cell(2, c + i).value) for i in range(6)]
        if vals == ["entrada", "salida", "ausente", "horas", "horas 50", "horas 100"]:
            bloques.append({
                "indice": len(bloques) + 1,
                "base_col": c - 1,
                "entrada": c,
                "salida": c + 1,
                "ausente": c + 2,
                "horas": c + 3,
                "horas_50": c + 4,
                "horas_100": c + 5,
            })
    if len(bloques) < 28:
        raise ValueError(f"No encontré suficientes bloques diarios. Detectados: {len(bloques)}")
    return bloques


def detectar_personal(ws, cols) -> list[dict]:
    personas = []
    fila_inicio = 3
    for r in range(fila_inicio, ws.max_row + 1):
        nombre = ws.cell(r, cols["col_nombre"]).value
        legajo = normalizar_legajo(ws.cell(r, cols["col_legajo"]).value)
        if nombre and normalizar_texto(nombre):
            personas.append({
                "fila": r,
                "legajo": legajo,
                "nombre": str(nombre).strip(),
                "centro": ws.cell(r, cols["col_centro"]).value if cols.get("col_centro") else None,
            })
    if not personas:
        raise ValueError("No encontré personal en la planilla.")
    return personas


# ---------------------------------------------------------------------------
# Feriados / fechas
# ---------------------------------------------------------------------------

DIAS_ES = {
    0: "lunes",
    1: "martes",
    2: "miércoles",
    3: "jueves",
    4: "viernes",
    5: "sábado",
    6: "domingo",
}
MESES_ES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
}


def formatear_fecha(fecha: date, es_feriado: bool = False) -> str:
    base = f"{DIAS_ES[fecha.weekday()]}, {fecha.day} de {MESES_ES[fecha.month]} de {fecha.year}"
    return f"FERIADO - {base}" if es_feriado else base


def feriados_argentina(anio: int, mes: int) -> dict[date, str]:
    kwargs = {"years": [anio], "language": "es"}
    if SUBDIV_FERIADOS:
        kwargs["subdiv"] = SUBDIV_FERIADOS
    cal = holidays.country_holidays(PAIS_FERIADOS, **kwargs)
    return {d: str(nombre) for d, nombre in cal.items() if d.year == anio and d.month == mes}


# ---------------------------------------------------------------------------
# Fórmulas diarias desde plantilla
# ---------------------------------------------------------------------------

def _formula_or_none(value):
    return value if isinstance(value, str) and value.startswith("=") else None


def extraer_plantillas_formula(ws, bloques, fila=3):
    """Toma fórmulas originales de la plantilla para laboral, sábado y domingo/feriado."""
    plantillas = {}
    for bloque in bloques:
        h = _formula_or_none(ws.cell(fila, bloque["horas"]).value)
        h50 = _formula_or_none(ws.cell(fila, bloque["horas_50"]).value)
        h100 = _formula_or_none(ws.cell(fila, bloque["horas_100"]).value)

        # Laboral: HORAS + HORAS 50%, sin 100%.
        if h and h50 and not h100 and "laboral" not in plantillas:
            plantillas["laboral"] = {
                "horas": (h, ws.cell(fila, bloque["horas"]).coordinate),
                "horas_50": (h50, ws.cell(fila, bloque["horas_50"]).coordinate),
                "horas_100": (None, None),
            }

        # Sábado: tiene HORAS y alguna fórmula de 100%; puede tener 50% propio de la plantilla.
        if h and h100 and "sabado" not in plantillas:
            plantillas["sabado"] = {
                "horas": (h, ws.cell(fila, bloque["horas"]).coordinate),
                "horas_50": (h50, ws.cell(fila, bloque["horas_50"]).coordinate) if h50 else (None, None),
                "horas_100": (h100, ws.cell(fila, bloque["horas_100"]).coordinate),
            }

        # Domingo/feriado: solo 100%.
        if h100 and not h and not h50 and "domingo" not in plantillas:
            plantillas["domingo"] = {
                "horas": (None, None),
                "horas_50": (None, None),
                "horas_100": (h100, ws.cell(fila, bloque["horas_100"]).coordinate),
            }

    faltan = [k for k in ["laboral", "sabado", "domingo"] if k not in plantillas]
    if faltan:
        raise ValueError(f"No pude extraer fórmulas base para: {', '.join(faltan)}")
    return plantillas


def traducir_formula(formula: str, origen: str, destino: str) -> str:
    if not formula:
        return None
    try:
        return Translator(formula, origin=origen).translate_formula(destino)
    except Exception:
        return formula


def aplicar_formula_diaria(ws, bloque, fila: int, tipo: str, plantillas):
    p = plantillas[tipo]
    for clave in ["horas", "horas_50", "horas_100"]:
        celda = ws.cell(fila, bloque[clave])
        formula, origen = p[clave]
        if formula and origen:
            celda.value = traducir_formula(formula, origen, celda.coordinate)
        else:
            celda.value = None


# ---------------------------------------------------------------------------
# Resúmenes y columnas adicionales
# ---------------------------------------------------------------------------

def _copiar_estilo_columna(ws, src_col, dst_col):
    for r in range(1, ws.max_row + 1):
        src = ws.cell(r, src_col)
        dst = ws.cell(r, dst_col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.font:
            dst.font = copy.copy(src.font)
        if src.border:
            dst.border = copy.copy(src.border)
    ws.column_dimensions[get_column_letter(dst_col)].width = ws.column_dimensions[get_column_letter(src_col)].width


def _find_header(ws, text, start_col=1):
    target = normalizar_header(text)
    for c in range(start_col, ws.max_column + 1):
        if normalizar_header(ws.cell(1, c).value) == target:
            return c
    return None


def _insert_after(ws, col, header):
    ws.insert_cols(col + 1)
    _copiar_estilo_columna(ws, col, col + 1)
    ws.cell(1, col + 1).value = header
    ws.cell(2, col + 1).value = None
    return col + 1


def asegurar_columnas_resumen(ws):
    """Asegura columnas ART, FERIADO y VACACIONES por quincena si la plantilla no las trae.

    Antes de insertar, descombina los encabezados de resumen de fila 1.
    Esto evita que una columna nueva quede dentro de un rango combinado y pierda su título.
    """
    starts = [
        _find_header(ws, "HORAS TOTALES") or _find_header(ws, "HORAS  TOTALES"),
        _find_header(ws, "HORAS PRIMERA QUINCENA"),
        _find_header(ws, "HORAS SEGUNDA QUINCENA"),
    ]
    starts = [c for c in starts if c]
    first_summary = min(starts) if starts else None
    if first_summary:
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row == 1 and rng.min_col >= first_summary:
                ws.unmerge_cells(str(rng))

    def find_zone_start(zone_name):
        if zone_name == "q1":
            return _find_header(ws, "HORAS PRIMERA QUINCENA")
        return _find_header(ws, "HORAS SEGUNDA QUINCENA")

    for zone_name in ["q1", "q2"]:
        start = find_zone_start(zone_name)
        if not start:
            continue

        def zone_end():
            next_hours = None
            for c in range(start + 1, ws.max_column + 1):
                h = normalizar_header(ws.cell(1, c).value)
                if zone_name == "q1" and "horas segunda quincena" in h:
                    next_hours = c
                    break
                if zone_name == "q2" and ("horas primera quincena" in h or "horas totales" in h):
                    next_hours = c
                    break
            return (next_hours - 1) if next_hours else ws.max_column

        def find_in_zone(name):
            end = zone_end()
            for c in range(start, end + 1):
                if normalizar_header(ws.cell(1, c).value) == normalizar_header(name):
                    return c
            return None

        vianda = find_in_zone("VIANDA")
        if vianda and find_in_zone("ART") is None:
            _insert_after(ws, vianda, "ART")

        premio = find_in_zone("PREMIO") or find_in_zone("ART") or find_in_zone("VIANDA")
        if premio and find_in_zone("FERIADO") is None:
            _insert_after(ws, premio, "FERIADO")

        feriado = find_in_zone("FERIADO") or premio
        if feriado and find_in_zone("VACACIONES") is None:
            _insert_after(ws, feriado, "VACACIONES")

    # Normalizar títulos de horas, ya sin celdas combinadas.
    for c, title in [
        (_find_header(ws, "HORAS TOTALES") or _find_header(ws, "HORAS  TOTALES"), "HORAS  TOTALES"),
        (_find_header(ws, "HORAS PRIMERA QUINCENA"), "HORAS PRIMERA QUINCENA"),
        (_find_header(ws, "HORAS SEGUNDA QUINCENA"), "HORAS  SEGUNDA QUINCENA"),
    ]:
        if c:
            ws.cell(1, c).value = title


def detectar_resumenes(ws, bloques, dias_mes):
    """Detecta columnas de horas y estado para totales/q1/q2 después de asegurar columnas."""
    q1_col = _find_header(ws, "HORAS PRIMERA QUINCENA")
    q2_col = _find_header(ws, "HORAS SEGUNDA QUINCENA")
    tot_col = _find_header(ws, "HORAS TOTALES") or _find_header(ws, "HORAS  TOTALES")
    if not q1_col or not q2_col or not tot_col:
        raise ValueError("No pude detectar columnas de HORAS TOTALES / PRIMERA / SEGUNDA QUINCENA.")

    def hours_block(start):
        return {
            "horas": start,
            "50": start + 1,
            "100": start + 2,
        }

    def extra_block(start, stop):
        result = {}
        targets = ["AUSENCIA", "ENFERMEDAD", "VIANDA", "ART", "FERIADO", "VACACIONES"]
        for c in range(start, stop + 1):
            h = normalizar_header(ws.cell(1, c).value)
            for t in targets:
                if h == normalizar_header(t):
                    result[t] = c
        return result

    q1_extra_start = q1_col + 3
    q1_extra_stop = q2_col - 1
    q2_extra_start = q2_col + 3
    q2_extra_stop = ws.max_column

    return {
        "totales": hours_block(tot_col),
        "q1": {**hours_block(q1_col), **extra_block(q1_extra_start, q1_extra_stop)},
        "q2": {**hours_block(q2_col), **extra_block(q2_extra_start, q2_extra_stop)},
    }


def formula_suma(celdas):
    return "=0" if not celdas else "=SUM(" + ",".join(celdas) + ")"


def formula_estado(estado: str, dias: list[tuple[date, dict]], fila: int, feriados: dict) -> str:
    terms = []
    for fecha, bloque in dias:
        if fecha.weekday() == 6 or fecha in feriados:
            continue
        peso = 4 if fecha.weekday() == 5 else 8
        ref = f"{get_column_letter(bloque['ausente'])}{fila}"
        terms.append(f'IF(UPPER(TRIM({ref}))="{estado}",{peso},0)')
    return "=0" if not terms else "=" + "+".join(terms)


def formula_vianda(dias: list[tuple[date, dict]], fila: int, feriados: dict) -> str:
    terms = []
    for fecha, bloque in dias:
        if fecha.weekday() == 6 or fecha in feriados:
            continue
        e = f"{get_column_letter(bloque['entrada'])}{fila}"
        s = f"{get_column_letter(bloque['salida'])}{fila}"
        a = f"{get_column_letter(bloque['ausente'])}{fila}"
        terms.append(f'IF(AND({e}<>"",{s}<>"",UPPER(TRIM({a}))<>"AUSENTE",UPPER(TRIM({a}))<>"ENFERMEDAD",UPPER(TRIM({a}))<>"ART",UPPER(TRIM({a}))<>"VACACIONES"),1,0)')
    return "=0" if not terms else "=" + "+".join(terms)


def formula_feriado(dias: list[tuple[date, dict]], feriados: dict) -> str:
    total = 0
    for fecha, _ in dias:
        if fecha in feriados and fecha.weekday() != 6:
            total += 4 if fecha.weekday() == 5 else 8
    return f"={total}"


def actualizar_resumenes(ws, personas, bloques, anio, mes, feriados):
    asegurar_columnas_resumen(ws)
    dias_mes = calendar.monthrange(anio, mes)[1]
    resumen = detectar_resumenes(ws, bloques, dias_mes)

    dias = []
    for d in range(1, min(dias_mes, len(bloques)) + 1):
        dias.append((date(anio, mes, d), bloques[d - 1]))
    dias_q1 = [x for x in dias if x[0].day <= 15]
    dias_q2 = [x for x in dias if x[0].day >= 16]

    for p in personas:
        fila = p["fila"]
        for qname, qdias in [("q1", dias_q1), ("q2", dias_q2)]:
            q = resumen[qname]
            ws.cell(fila, q["horas"]).value = formula_suma([f"{get_column_letter(b['horas'])}{fila}" for _, b in qdias])
            ws.cell(fila, q["50"]).value = formula_suma([f"{get_column_letter(b['horas_50'])}{fila}" for _, b in qdias])
            ws.cell(fila, q["100"]).value = formula_suma([f"{get_column_letter(b['horas_100'])}{fila}" for _, b in qdias])

            if "AUSENCIA" in q:
                ws.cell(fila, q["AUSENCIA"]).value = formula_estado("AUSENTE", qdias, fila, feriados)
            if "ENFERMEDAD" in q:
                ws.cell(fila, q["ENFERMEDAD"]).value = formula_estado("ENFERMEDAD", qdias, fila, feriados)
            if "ART" in q:
                ws.cell(fila, q["ART"]).value = formula_estado("ART", qdias, fila, feriados)
            if "VACACIONES" in q:
                ws.cell(fila, q["VACACIONES"]).value = formula_estado("VACACIONES", qdias, fila, feriados)
            if "VIANDA" in q:
                ws.cell(fila, q["VIANDA"]).value = formula_vianda(qdias, fila, feriados)
            if "FERIADO" in q:
                ws.cell(fila, q["FERIADO"]).value = formula_feriado(qdias, feriados)

        # Totales de horas.
        t = resumen["totales"]
        q1 = resumen["q1"]
        q2 = resumen["q2"]
        ws.cell(fila, t["horas"]).value = f"={get_column_letter(q1['horas'])}{fila}+{get_column_letter(q2['horas'])}{fila}"
        ws.cell(fila, t["50"]).value = f"={get_column_letter(q1['50'])}{fila}+{get_column_letter(q2['50'])}{fila}"
        ws.cell(fila, t["100"]).value = f"={get_column_letter(q1['100'])}{fila}+{get_column_letter(q2['100'])}{fila}"

    # Corregir formato y color de fuente en todas las celdas numéricas del resumen.
    normalizar_formato_resumen(ws, personas, resumen)
    return resumen


# ---------------------------------------------------------------------------
# Normalización visual de resúmenes y alertas
# ---------------------------------------------------------------------------

def _poner_texto_negro(celda):
    """Conserva tipografía/tamaño, pero fuerza el color negro."""
    fuente = copy.copy(celda.font)
    fuente.color = COLOR_TEXTO_RESUMEN
    celda.font = fuente


def normalizar_formato_resumen(ws, personas, resumen):
    """Evita resultados invisibles o mostrados como horas.

    Las columnas insertadas pueden heredar:
    - formatos de hora como [h]:mm:ss;
    - colores de fuente iguales al fondo;
    - formatos personalizados que ocultan el cero.

    Todas las celdas numéricas de resumen se fijan como enteros visibles.
    """
    columnas = set()

    for zona in ("totales", "q1", "q2"):
        for clave, columna in resumen[zona].items():
            if isinstance(columna, int):
                columnas.add(columna)

    for persona in personas:
        fila = persona["fila"]
        for columna in columnas:
            celda = ws.cell(fila, columna)
            celda.number_format = "0"
            _poner_texto_negro(celda)


def _fill_rojo_faltas(cantidad):
    for minimo, color, etiqueta in ESCALA_ROJO_FALTAS:
        if cantidad >= minimo:
            return PatternFill("solid", fgColor=color), etiqueta
    return None, None


def _limpiar_colores_alerta(ws, personas, cols):
    """Borra alertas de ejecuciones anteriores únicamente en LEG y NOMBRE."""
    sin_relleno = PatternFill(fill_type=None)
    for persona in personas:
        fila = persona["fila"]
        ws.cell(fila, cols["col_legajo"]).fill = copy.copy(sin_relleno)
        ws.cell(fila, cols["col_nombre"]).fill = copy.copy(sin_relleno)


def _razones_vinculo_dudoso(personas, cargados, ignorados):
    """Devuelve motivos de revisión agrupados por fila de la planilla."""
    razones = defaultdict(list)

    cantidad_por_legajo = defaultdict(int)
    persona_por_legajo = defaultdict(list)
    persona_por_nombre = defaultdict(list)

    for persona in personas:
        legajo = normalizar_legajo(persona.get("legajo"))
        if legajo:
            cantidad_por_legajo[legajo] += 1
            persona_por_legajo[legajo].append(persona)
        persona_por_nombre[clave_nombre(persona.get("nombre"))].append(persona)

    for item in cargados:
        fila = item.get("FILA")
        if not fila:
            continue

        metodo = str(item.get("METODO") or "").strip().lower()
        try:
            puntaje = float(item.get("PUNTAJE") or 0)
        except Exception:
            puntaje = 0.0

        leg_fichador = normalizar_legajo(item.get("LEG_FICHADOR"))
        leg_planilla = normalizar_legajo(item.get("LEG_PLANILLA"))

        if metodo != "legajo+nombre":
            razones[fila].append(f"Método de vínculo: {item.get('METODO')}")
        if puntaje < 0.999:
            razones[fila].append(f"Coincidencia de nombre parcial ({puntaje:.3f})")
        if leg_fichador and leg_planilla and leg_fichador != leg_planilla:
            razones[fila].append(
                f"Legajo fichador {leg_fichador} distinto de planilla {leg_planilla}"
            )
        if leg_planilla and cantidad_por_legajo.get(leg_planilla, 0) > 1:
            razones[fila].append(f"Legajo {leg_planilla} repetido en la planilla")

    # También advertir cuando un fichaje ignorado quedó cerca de una persona.
    for item in ignorados:
        try:
            puntaje = float(item.get("PUNTAJE") or 0)
        except Exception:
            puntaje = 0.0

        if puntaje < UMBRAL_IGNORADO_CERCANO:
            continue

        leg_candidato = normalizar_legajo(item.get("LEG_CANDIDATO"))
        candidatos = persona_por_legajo.get(leg_candidato, []) if leg_candidato else []

        if not candidatos:
            candidatos = persona_por_nombre.get(
                clave_nombre(item.get("MEJOR_CANDIDATO")),
                [],
            )

        if len(candidatos) == 1:
            persona = candidatos[0]
            razones[persona["fila"]].append(
                f"Fichaje ignorado cercano: {item.get('NOMBRE_FICHADOR')} "
                f"({puntaje:.3f})"
            )

    # Eliminar razones repetidas manteniendo orden.
    return {
        fila: list(dict.fromkeys(motivos))
        for fila, motivos in razones.items()
    }


def aplicar_alertas_visuales(
    ws,
    personas,
    bloques,
    cols,
    anio,
    mes,
    max_fecha,
    cargados,
    ignorados,
):
    """Colorea únicamente:
    - NOMBRE: escala roja según cantidad de días AUSENTE.
    - LEG: amarillo si el emparejamiento requiere revisión.
    """
    _limpiar_colores_alerta(ws, personas, cols)
    razones_dudosas = _razones_vinculo_dudoso(personas, cargados, ignorados)
    alertas = []

    dias_mes = calendar.monthrange(anio, mes)[1]

    for persona in personas:
        fila = persona["fila"]
        fechas_ausente = []

        for dia in range(1, min(dias_mes, len(bloques)) + 1):
            fecha = date(anio, mes, dia)
            if fecha > max_fecha:
                continue
            bloque = bloques[dia - 1]
            estado = str(ws.cell(fila, bloque["ausente"]).value or "").strip().upper()
            if estado == "AUSENTE":
                fechas_ausente.append(fecha)

        relleno_rojo, nivel = _fill_rojo_faltas(len(fechas_ausente))
        if relleno_rojo is not None:
            ws.cell(fila, cols["col_nombre"]).fill = copy.copy(relleno_rojo)
            alertas.append({
                "TIPO": "AUSENCIAS",
                "COLOR": nivel,
                "FILA": fila,
                "LEG": persona.get("legajo"),
                "PERSONA": persona.get("nombre"),
                "CANTIDAD": len(fechas_ausente),
                "DETALLE": ", ".join(
                    fecha.strftime("%d/%m/%Y") for fecha in fechas_ausente
                ),
                "CRITERIO": f"{len(fechas_ausente)} días marcados AUSENTE",
            })

        motivos = razones_dudosas.get(fila, [])
        if motivos:
            ws.cell(fila, cols["col_legajo"]).fill = PatternFill(
                "solid",
                fgColor=COLOR_AMARILLO_DUDOSO,
            )
            alertas.append({
                "TIPO": "REVISAR VÍNCULO",
                "COLOR": "AMARILLO",
                "FILA": fila,
                "LEG": persona.get("legajo"),
                "PERSONA": persona.get("nombre"),
                "CANTIDAD": len(motivos),
                "DETALLE": " | ".join(motivos),
                "CRITERIO": "Posible confusión de nombre o legajo",
            })

    return alertas

# ---------------------------------------------------------------------------
# Matching y carga
# ---------------------------------------------------------------------------

def vincular_personas(personas: list[dict], registros: list[dict]):
    por_legajo = defaultdict(list)
    por_clave = defaultdict(list)
    for p in personas:
        if p.get("legajo"):
            por_legajo[p["legajo"]].append(p)
        por_clave[clave_nombre(p["nombre"])].append(p)

    identidades = sorted({(r["id_raw"], r["legajo"], r["nombre"]) for r in registros})
    usados = set()
    vinculos = {}
    cargados = []
    ignorados = []

    for raw_id, legajo, nombre in identidades:
        elegido = None
        candidatos = []
        for p in por_legajo.get(legajo, []):
            score = puntaje_nombre(nombre, p["nombre"])
            if nombres_compatibles(nombre, p["nombre"]):
                candidatos.append((score, p, "legajo+nombre"))
        candidatos.sort(key=lambda x: x[0], reverse=True)
        if candidatos and (len(candidatos) == 1 or candidatos[0][0] - candidatos[1][0] >= 0.03):
            elegido = candidatos[0]

        if elegido is None:
            exactos = [p for p in por_clave.get(clave_nombre(nombre), []) if p["fila"] not in usados]
            if len(exactos) == 1:
                elegido = (1.0, exactos[0], "nombre exacto/reordenado")

        if elegido is None:
            ranking = sorted(
                [(puntaje_nombre(nombre, p["nombre"]), p) for p in personas if p["fila"] not in usados],
                key=lambda x: x[0], reverse=True
            )
            if ranking and ranking[0][0] >= UMBRAL_NOMBRE_SIMILAR and (len(ranking) == 1 or ranking[0][0] - ranking[1][0] >= 0.06):
                elegido = (ranking[0][0], ranking[0][1], "nombre similar")

        if elegido is not None:
            score, persona, metodo = elegido
            if persona["fila"] not in usados:
                vinculos[(raw_id, nombre)] = persona
                usados.add(persona["fila"])
                cargados.append({
                    "ID_FICHADOR": raw_id,
                    "LEG_FICHADOR": legajo,
                    "NOMBRE_FICHADOR": nombre,
                    "LEG_PLANILLA": persona.get("legajo"),
                    "NOMBRE_PLANILLA": persona.get("nombre"),
                    "FILA": persona.get("fila"),
                    "METODO": metodo,
                    "PUNTAJE": round(score, 3),
                })
                continue

        mejor = sorted([(puntaje_nombre(nombre, p["nombre"]), p) for p in personas], key=lambda x: x[0], reverse=True)
        mejor_score, mejor_p = mejor[0] if mejor else (None, None)
        ignorados.append({
            "ID_FICHADOR": raw_id,
            "LEG_FICHADOR": legajo,
            "NOMBRE_FICHADOR": nombre,
            "MEJOR_CANDIDATO": mejor_p["nombre"] if mejor_p else "",
            "LEG_CANDIDATO": mejor_p["legajo"] if mejor_p else "",
            "PUNTAJE": round(mejor_score, 3) if mejor_score is not None else "",
            "MOTIVO": "No pertenece a la dotación o no tiene coincidencia segura.",
        })

    return vinculos, cargados, ignorados


def agrupar_marcas(registros, vinculos):
    marcas = defaultdict(list)
    for r in registros:
        persona = vinculos.get((r["id_raw"], r["nombre"]))
        if persona:
            marcas[(persona["fila"], r["dt"].date())].append(r["dt"])
    return marcas


def limpiar_auxiliares(wb):
    for h in HOJAS_AUXILIARES:
        if h in wb.sheetnames:
            del wb[h]


def agregar_hoja_control(wb, nombre, rows: list[dict]):
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)
    if not rows:
        ws.append(["SIN REGISTROS"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col_cells in ws.columns:
        letter = col_cells[0].column_letter
        width = max(12, min(45, max(len(str(c.value or "")) for c in col_cells) + 2))
        ws.column_dimensions[letter].width = width


# ---------------------------------------------------------------------------
# Detección de archivos y proceso principal
# ---------------------------------------------------------------------------

def score_planilla(path: str | Path) -> int:
    if Path(path).suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return 0
    score = 0
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        for ws in wb.worksheets[:3]:
            texto = []
            for r in range(1, min(ws.max_row, 5) + 1):
                for c in range(1, min(ws.max_column, 260) + 1):
                    texto.append(str(ws.cell(r, c).value or ""))
            flat = normalizar_texto(" ".join(texto))
            if "apellido y nombre" in flat:
                score += 4
            if "entrada" in flat and "salida" in flat and "ausente" in flat:
                score += 4
            if "horas primera quincena" in flat:
                score += 2
        wb.close()
    except Exception:
        return 0
    return score


def score_fichajes(path: str | Path) -> int:
    if Path(path).suffix.lower() != ".xls":
        return 0
    try:
        libro = xlrd.open_workbook(str(path))
        sh = libro.sheet_by_index(0)
        headers = " ".join(normalizar_header(sh.cell_value(0, c)) for c in range(sh.ncols))
        if "id" in headers and "nombre" in headers and "hora" in headers:
            return 10
    except Exception:
        return 0
    return 0


def detectar_archivos(archivos):
    diag = []
    for p in archivos:
        if Path(p).suffix.lower() not in {".xls", ".xlsx", ".xlsm", ".xltx", ".xltm"}:
            continue
        diag.append({
            "path": str(p),
            "score_planilla": score_planilla(p),
            "score_fichajes": score_fichajes(p),
        })
    if len(diag) < 2:
        raise ValueError("Subí al menos una planilla base y un archivo de fichajes.")
    planilla = max(diag, key=lambda x: x["score_planilla"])
    fichajes = max(diag, key=lambda x: x["score_fichajes"])
    if planilla["score_planilla"] <= 0 or fichajes["score_fichajes"] <= 0 or planilla["path"] == fichajes["path"]:
        raise ValueError("No pude detectar automáticamente planilla base y fichajes.")
    return planilla["path"], fichajes["path"], diag


def _set_calc_mode(wb):
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def procesar_archivos(archivo_planilla, archivo_fichajes, descargar_en_colab=False):
    registros = leer_fichajes(archivo_fichajes)
    anio, mes = detectar_mes(registros)
    dias_mes = calendar.monthrange(anio, mes)[1]
    max_fecha = max(r["dt"].date() for r in registros)
    min_fecha = min(r["dt"].date() for r in registros)

    wb = load_workbook(archivo_planilla, data_only=False)
    ws = detectar_hoja_planilla(wb)
    limpiar_auxiliares(wb)
    _set_calc_mode(wb)

    cols = detectar_columnas_identidad(ws)
    bloques = detectar_bloques_diarios(ws)
    personas = detectar_personal(ws, cols)
    plantillas = extraer_plantillas_formula(ws, bloques, fila=personas[0]["fila"])
    feriados = feriados_argentina(anio, mes)

    if len(bloques) < dias_mes:
        raise ValueError(f"La planilla tiene {len(bloques)} bloques diarios, pero el mes tiene {dias_mes} días.")

    vinculos, cargados, ignorados = vincular_personas(personas, registros)
    marcas = agrupar_marcas(registros, vinculos)

    # Reconfigurar días, fórmulas y fichadas.
    for i, bloque in enumerate(bloques, start=1):
        visible = i <= dias_mes
        fecha = date(anio, mes, i) if visible else None
        if visible:
            es_feriado = fecha in feriados
            ws.cell(1, bloque["base_col"]).value = formatear_fecha(fecha, es_feriado)
        else:
            ws.cell(1, bloque["base_col"]).value = None

        for c in range(bloque["base_col"], bloque["horas_100"] + 1):
            ws.column_dimensions[get_column_letter(c)].hidden = not visible

        for p in personas:
            fila = p["fila"]
            for k in ["entrada", "salida", "ausente"]:
                ws.cell(fila, bloque[k]).value = None

            if visible:
                if fecha in feriados or fecha.weekday() == 6:
                    tipo = "domingo"
                elif fecha.weekday() == 5:
                    tipo = "sabado"
                else:
                    tipo = "laboral"
                aplicar_formula_diaria(ws, bloque, fila, tipo, plantillas)

                lista = sorted(marcas.get((fila, fecha), []))
                if lista:
                    entrada = lista[0].time().replace(second=0, microsecond=0)
                    salida = lista[-1].time().replace(second=0, microsecond=0) if len(lista) >= 2 else None
                    ws.cell(fila, bloque["entrada"]).value = entrada
                    ws.cell(fila, bloque["entrada"]).number_format = "hh:mm"
                    if salida and salida != entrada:
                        ws.cell(fila, bloque["salida"]).value = salida
                        ws.cell(fila, bloque["salida"]).number_format = "hh:mm"
                else:
                    if fecha <= max_fecha and fecha.weekday() != 6 and fecha not in feriados:
                        ws.cell(fila, bloque["ausente"]).value = "AUSENTE"
            else:
                for k in ["horas", "horas_50", "horas_100"]:
                    ws.cell(fila, bloque[k]).value = None

    # Resúmenes de horas y nuevos estados.
    resumen = actualizar_resumenes(ws, personas, bloques, anio, mes, feriados)

    # Alertas recalculadas desde cero.
    alertas = aplicar_alertas_visuales(
        ws=ws,
        personas=personas,
        bloques=bloques,
        cols=cols,
        anio=anio,
        mes=mes,
        max_fecha=max_fecha,
        cargados=cargados,
        ignorados=ignorados,
    )

    agregar_hoja_control(wb, "VINCULOS_CARGADOS", cargados)
    agregar_hoja_control(wb, "IGNORADOS_FICHADOR", ignorados)
    agregar_hoja_control(wb, "ALERTAS_PROCESO", alertas)
    agregar_hoja_control(wb, "DIAGNOSTICO", [{
        "PLANILLA": Path(archivo_planilla).name,
        "FICHAJES": Path(archivo_fichajes).name,
        "MES_DETECTADO": f"{anio}-{mes:02d}",
        "FECHA_MIN_FICHAJES": min_fecha.isoformat(),
        "FECHA_MAX_FICHAJES": max_fecha.isoformat(),
        "PERSONAS_PLANILLA": len(personas),
        "IDENTIDADES_VINCULADAS": len(cargados),
        "IDENTIDADES_IGNORADAS": len(ignorados),
        "NOMBRES_RESALTADOS_ROJO": sum(1 for a in alertas if a.get("TIPO") == "AUSENCIAS"),
        "LEGAJOS_RESALTADOS_AMARILLO": sum(1 for a in alertas if a.get("TIPO") == "REVISAR VÍNCULO"),
        "FERIADOS_DETECTADOS": "; ".join(f"{d.strftime('%d/%m/%Y')} {feriados[d]}" for d in sorted(feriados)),
    }])

    nombre_base = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(archivo_planilla).stem).strip("_")
    salida = f"PLANILLA_COMPLETA_{nombre_base}_{anio}_{mes:02d}.xlsx"
    wb.save(salida)
    return [salida]


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Procesar planilla INMAC.")
    parser.add_argument("archivos", nargs="+", help="Planilla base y archivo de fichajes")
    args = parser.parse_args()
    planilla, fichajes, diag = detectar_archivos(args.archivos)
    print("Planilla:", planilla)
    print("Fichajes:", fichajes)
    print(procesar_archivos(planilla, fichajes, descargar_en_colab=False))

